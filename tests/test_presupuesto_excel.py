"""Presupuesto Excel Expertis: lectura, alta de obra y contraste con el PDF del cliente."""

from decimal import Decimal
from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook
from openpyxl.styles import Font
from sqlalchemy.orm import Session

from app.db import SessionLocal
from app.main import app
from app.models import Nave, Presupuesto, Proyecto, Tarea
from app.schemas.extraccion import ApartadoExtraido, NaveExtraida, PresupuestoExtraido
from app.services.excel_expertis import leer_excel_expertis


@pytest.fixture
def client():
    with TestClient(app) as c:
        yield c


@pytest.fixture
def db() -> Session:
    session = SessionLocal()
    try:
        yield session
    finally:
        session.close()


def _libro_descuadre() -> bytes:
    wb = Workbook()
    hoja = wb.active
    hoja.title = "VENTILACION"
    hoja["B4"] = "Cliente"
    hoja["B8"] = "SISTEMA DE VENTILACION FICTICIO"
    cabeceras = {
        "B10": "Cant.",
        "C10": "Cod.",
        "D10": "Descripción",
        "H10": "Ref. (1)",
        "I10": "Ref. (2)",
        "J10": "Euros",
        "K10": "Dto. %",
        "L10": "TOTAL",
    }
    for celda, texto in cabeceras.items():
        hoja[celda] = texto
    hoja["D11"] = "SALIDA DE SILOS"
    hoja["D11"].font = Font(italic=True)
    hoja["B12"] = 0
    hoja["C12"] = "07"
    hoja["D12"] = "CANTIDAD CERO"
    hoja["J12"] = 99
    hoja["L12"] = 99
    hoja["B13"] = 2
    hoja["C13"] = "08"
    hoja["D13"] = "TRANSPORTADOR"
    hoja["H13"] = "REF-1"
    hoja["I13"] = "REF-2"
    hoja["J13"] = 10
    hoja["K13"] = 0
    hoja["L13"] = 20
    hoja["D14"] = "total silos"
    hoja["L14"] = 20
    hoja["D15"] = "EXAFEED"
    hoja["D15"].font = Font(italic=True)
    hoja["B16"] = 1
    hoja["C16"] = "08"
    hoja["D16"] = "ELEMENTOS DE MONTAJE"
    hoja["J16"] = 5
    hoja["L16"] = 5
    hoja["B17"] = 1
    hoja["C17"] = "58"
    hoja["D17"] = "MONTAJE TRANSPORTADORES"
    hoja["J17"] = 30
    hoja["L17"] = 30
    hoja["D18"] = "total general"
    hoja["L18"] = 40
    hoja["B19"] = "Detalles de pago"
    hoja["B20"] = 9
    hoja["C20"] = "99"
    hoja["D20"] = "NO DEBE ENTRAR"
    hoja["L20"] = 999
    datos = wb.create_sheet("DATA")
    datos["A1"] = "catalogo interno"
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _libro_formula() -> bytes:
    wb = Workbook()
    hoja = wb.active
    hoja.title = "Presupuesto"
    for col, texto in {
        "B": "Cant.",
        "C": "Cod.",
        "D": "Descripción",
        "J": "Euros",
        "K": "Dto. %",
        "L": "TOTAL",
    }.items():
        hoja[f"{col}1"] = texto
    hoja["B2"] = 2
    hoja["C2"] = "11"
    hoja["D2"] = "PANEL"
    hoja["J2"] = 10
    hoja["K2"] = 10
    hoja["L2"] = "=B2*J2*(1-K2/100)"
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_lee_montaje_y_avisa_si_el_total_no_cuadra():
    leido = leer_excel_expertis(_libro_descuadre())

    assert leido.nombre_apartado == "SISTEMA DE VENTILACION FICTICIO"
    assert leido.partidas == 3
    assert leido.lineas_montaje == 1
    assert leido.sumas_cuadran is False
    assert leido.avisos[0].bloque == "total general"
    assert leido.avisos[0].suma_lineas == Decimal("55.00")
    assert leido.avisos[0].total_documento == Decimal("40.00")
    assert leido.importe == Decimal("55.00")

    silos = leido.nodos[0]
    assert silos.nivel == "subapartado"
    assert silos.descripcion == "SALIDA DE SILOS"
    material = silos.hijos[0]
    assert material.codigo == "08"
    assert material.referencia == "REF-1"
    assert material.referencia_alt == "REF-2"
    assert material.es_montaje is False
    assert material.importe == Decimal("20.00")
    assert len(silos.hijos) == 1
    assert silos.importe == Decimal("20.00")

    exafeed = leido.nodos[1]
    assert exafeed.descripcion == "EXAFEED"
    montaje = exafeed.hijos[0]
    assert montaje.nivel == "subapartado"
    assert montaje.descripcion == "Montaje EXAFEED"
    assert [hijo.descripcion for hijo in montaje.hijos] == [
        "MONTAJE TRANSPORTADORES",
    ]
    assert all(hijo.es_montaje for hijo in montaje.hijos)
    assert [hijo.capitulo for hijo in montaje.hijos] == ["EXAFEED"]
    assert montaje.importe == Decimal("30.00")
    elementos = exafeed.hijos[1]
    assert elementos.descripcion == "ELEMENTOS DE MONTAJE"
    assert elementos.es_montaje is False
    assert elementos.importe == Decimal("5.00")
    assert exafeed.importe == Decimal("35.00")
    descripciones = [silos.descripcion, exafeed.descripcion, montaje.descripcion] + [
        hijo.descripcion for hijo in montaje.hijos + silos.hijos + exafeed.hijos
    ]
    assert "NO DEBE ENTRAR" not in descripciones
    assert "CANTIDAD CERO" not in descripciones


def _libro_subapartados() -> bytes:
    wb = Workbook()
    hoja = wb.active
    hoja.title = "ALIMENTACION"
    for col, texto in {
        "B": "Cant.",
        "C": "Cod.",
        "D": "Descripción",
        "J": "Euros",
        "K": "Dto. %",
        "L": "TOTAL",
    }.items():
        hoja[f"{col}1"] = texto
    negrita = Font(bold=True)
    hoja["C2"] = "07"
    hoja["D2"] = "SILOS FICTICIO"
    hoja["C2"].font = negrita
    hoja["D2"].font = negrita
    hoja["B3"] = 2
    hoja["C3"] = "07"
    hoja["D3"] = "PIEZA DE SILO"
    hoja["J3"] = 10
    hoja["L3"] = 20
    hoja["C4"] = "08"
    hoja["D4"] = "LINEA FICTICIA"
    hoja["C4"].font = negrita
    hoja["D4"].font = negrita
    hoja["D5"] = "SALIDA FICTICIA"
    hoja["D5"].font = Font(italic=True)
    hoja["B6"] = 1
    hoja["C6"] = "08"
    hoja["D6"] = "CONO FICTICIO"
    hoja["J6"] = 5
    hoja["K6"] = 10
    hoja["L6"] = Decimal("4.50")
    hoja["B7"] = 1
    hoja["C7"] = "08"
    hoja["D7"] = "MONTAJE CONO"
    hoja["J7"] = 3
    hoja["L7"] = 3
    hoja["D8"] = "REPARTO FICTICIO"
    hoja["D8"].font = Font(italic=True)
    hoja["B9"] = 1
    hoja["C9"] = "08"
    hoja["D9"] = "MAQUINA FICTICIA"
    hoja["J9"] = 4
    hoja["L9"] = 4
    hoja["D10"] = "total general"
    hoja["L10"] = Decimal("31.50")
    hoja["C11"] = "11"
    hoja["D11"] = "NOTA SIN NEGRITA"
    hoja["C12"] = "11"
    hoja["D12"] = "REJILLA CON PRECIO"
    hoja["J12"] = 10
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_los_codigos_sin_cantidad_abren_subapartados():
    leido = leer_excel_expertis(_libro_subapartados())

    assert leido.partidas == 4
    assert leido.sumas_cuadran is True

    silos = leido.nodos[0]
    assert silos.codigo == "07"
    assert silos.nivel == "subapartado"
    assert silos.descripcion == "SILOS FICTICIO"
    assert [hijo.descripcion for hijo in silos.hijos] == ["PIEZA DE SILO"]
    assert silos.importe == Decimal("20.00")

    linea = leido.nodos[1]
    assert linea.codigo == "08"
    assert linea.descripcion == "LINEA FICTICIA"
    assert [hijo.descripcion for hijo in linea.hijos] == ["SALIDA FICTICIA", "REPARTO FICTICIO"]
    salida = linea.hijos[0]
    assert salida.nivel == "subapartado"
    montaje = salida.hijos[0]
    assert montaje.descripcion == "Montaje SALIDA FICTICIA"
    assert [hijo.descripcion for hijo in montaje.hijos] == ["MONTAJE CONO"]
    assert montaje.hijos[0].capitulo == "SALIDA FICTICIA"
    assert [hijo.descripcion for hijo in salida.hijos] == [
        "Montaje SALIDA FICTICIA",
        "CONO FICTICIO",
    ]
    cono = salida.hijos[1]
    assert cono.cantidad == Decimal("1")
    assert cono.descuento_pct == Decimal("10")
    reparto = linea.hijos[1]
    assert [hijo.descripcion for hijo in reparto.hijos] == ["MAQUINA FICTICIA"]
    assert linea.importe == Decimal("11.50")
    descripciones = [nodo.descripcion for nodo in leido.nodos]
    descripciones += [hijo.descripcion for nodo in leido.nodos for hijo in nodo.hijos]
    assert "NOTA SIN NEGRITA" not in descripciones
    assert "REJILLA CON PRECIO" not in descripciones


def _libro_montaje_general() -> bytes:
    wb = Workbook()
    hoja = wb.active
    hoja.title = "ALIMENTACION"
    for col, texto in {
        "B": "Cant.",
        "C": "Cod.",
        "D": "Descripción",
        "J": "Euros",
        "L": "TOTAL",
    }.items():
        hoja[f"{col}1"] = texto
    negrita = Font(bold=True)
    hoja["B2"] = 1
    hoja["C2"] = "01"
    hoja["D2"] = "MONTAJE VARIOS BLOQUES"
    hoja["J2"] = 8
    hoja["L2"] = 8
    hoja["C3"] = "07"
    hoja["D3"] = "CIKKIS"
    hoja["C3"].font = negrita
    hoja["D3"].font = negrita
    hoja["B4"] = 1
    hoja["C4"] = "07"
    hoja["D4"] = "PIEZA CIKKIS"
    hoja["J4"] = 10
    hoja["L4"] = 10
    hoja["B5"] = 1
    hoja["C5"] = "07"
    hoja["D5"] = "MONTAJE CIKKIS"
    hoja["J5"] = 4
    hoja["L5"] = 4
    hoja["C6"] = "08"
    hoja["D6"] = "TOLVA FICTICIA"
    hoja["C6"].font = negrita
    hoja["D6"].font = negrita
    hoja["B7"] = 1
    hoja["C7"] = "08"
    hoja["D7"] = "PIEZA TOLVA"
    hoja["J7"] = 6
    hoja["L7"] = 6
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_el_montaje_de_un_subapartado_lleva_su_nombre_y_el_comun_es_general():
    leido = leer_excel_expertis(_libro_montaje_general())

    assert [nodo.descripcion for nodo in leido.nodos] == [
        "Montaje general",
        "CIKKIS",
        "TOLVA FICTICIA",
    ]
    general = leido.nodos[0]
    assert [hijo.descripcion for hijo in general.hijos] == ["MONTAJE VARIOS BLOQUES"]
    assert general.importe == Decimal("8.00")

    cikkis = leido.nodos[1]
    assert [hijo.descripcion for hijo in cikkis.hijos] == ["Montaje CIKKIS", "PIEZA CIKKIS"]
    assert [hijo.descripcion for hijo in cikkis.hijos[0].hijos] == ["MONTAJE CIKKIS"]
    assert cikkis.hijos[0].hijos[0].capitulo == "CIKKIS"
    assert cikkis.importe == Decimal("14.00")
    assert leido.nodos[2].importe == Decimal("6.00")
    assert leido.importe == Decimal("28.00")


def _libro_montaje_al_final() -> bytes:
    """El montaje va junto al final, en GENERAL, como en ventilación. Importes ficticios."""
    wb = Workbook()
    hoja = wb.active
    hoja.title = "VENTILACION"
    for col, texto in {
        "B": "Cant.",
        "C": "Cod.",
        "D": "Descripción",
        "J": "Euros",
        "L": "TOTAL",
    }.items():
        hoja[f"{col}1"] = texto
    negrita = Font(bold=True)
    filas = [
        (2, "07", "VENTANAS FICTICIAS", None, None),
        (3, "07", "PIEZA DE VENTANA", 2, 2),
        (4, "08", "CHIMENEAS FICTICIAS", None, None),
        (5, "08", "PIEZA DE CHIMENEA", 3, 3),
        (6, "09", "CIKKIS", None, None),
        (7, "09", "PIEZA CIKKIS", 4, 4),
        (8, "50", "GENERAL", None, None),
        (9, "50", "MONTAJE VENTANAS FICTICIAS", 5, 5),
        (10, "50", "MONTAJE TUBOS CHIMENEAS FICTICIAS", 6, 6),
        (11, "50", "MONTAJE CIKKIS", 7, 7),
        (12, "50", "MONTAJE DIETAS FICTICIAS", 8, 8),
        (13, "50", "TRANSPORTE FICTICIO", 9, 9),
    ]
    for fila, codigo, descripcion, euros, total in filas:
        hoja[f"C{fila}"] = codigo
        hoja[f"D{fila}"] = descripcion
        if euros is None:
            hoja[f"C{fila}"].font = negrita
            hoja[f"D{fila}"].font = negrita
            continue
        hoja[f"B{fila}"] = 1
        hoja[f"J{fila}"] = euros
        hoja[f"L{fila}"] = total
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_el_montaje_de_general_se_relaciona_con_el_producto():
    from app.services.excel_expertis import (
        aplicar_relacion_montaje,
        asignar_montaje_por_nombre,
        lineas_y_productos_montaje,
    )

    leido = leer_excel_expertis(_libro_montaje_al_final())
    lineas, productos = lineas_y_productos_montaje(leido.nodos)
    assert [nodo.descripcion for nodo in productos] == [
        "VENTANAS FICTICIAS",
        "CHIMENEAS FICTICIAS",
        "CIKKIS",
    ]
    assert [nodo.descripcion for nodo in lineas] == [
        "MONTAJE VENTANAS FICTICIAS",
        "MONTAJE TUBOS CHIMENEAS FICTICIAS",
        "MONTAJE CIKKIS",
        "MONTAJE DIETAS FICTICIAS",
    ]

    asignaciones = asignar_montaje_por_nombre(
        [nodo.descripcion for nodo in lineas],
        [nodo.descripcion for nodo in productos],
    )
    assert asignaciones == [0, 1, 2, None]
    leido.importe = aplicar_relacion_montaje(leido.nodos, asignaciones)

    ventanas, chimeneas, cikkis, general_resto, montaje_general = leido.nodos
    assert ventanas.hijos[0].descripcion == "Montaje VENTANAS FICTICIAS"
    assert [hijo.descripcion for hijo in ventanas.hijos[0].hijos] == ["MONTAJE VENTANAS FICTICIAS"]
    assert chimeneas.hijos[0].descripcion == "Montaje CHIMENEAS FICTICIAS"
    assert cikkis.hijos[0].descripcion == "Montaje CIKKIS"
    assert [hijo.descripcion for hijo in general_resto.hijos] == ["TRANSPORTE FICTICIO"]
    assert montaje_general.descripcion == "Montaje general"
    assert [hijo.descripcion for hijo in montaje_general.hijos] == ["MONTAJE DIETAS FICTICIAS"]
    assert leido.importe == Decimal("44.00")


def test_la_relacion_de_montaje_usa_el_llm(monkeypatch):
    from app.services.presupuesto_excel import relacionar_montaje_general

    leido = leer_excel_expertis(_libro_montaje_al_final())

    def hay_llm(_settings):
        return True

    def completar_texto(**_kwargs):
        return (
            '{"asignaciones": ['
            '{"linea": 0, "producto": 0},'
            '{"linea": 1, "producto": "CHIMENEAS FICTICIAS"},'
            '{"linea": 2, "producto": 2},'
            '{"linea": 3, "producto": null}'
            "]}"
        )

    monkeypatch.setattr("app.services.presupuesto_excel.llm_svc.hay_llm", hay_llm)
    monkeypatch.setattr("app.services.presupuesto_excel.llm_svc.completar_texto", completar_texto)
    relacionar_montaje_general(leido)

    assert leido.nodos[0].hijos[0].descripcion == "Montaje VENTANAS FICTICIAS"
    assert leido.nodos[1].hijos[0].descripcion == "Montaje CHIMENEAS FICTICIAS"
    assert leido.nodos[2].hijos[0].descripcion == "Montaje CIKKIS"
    assert leido.nodos[-1].descripcion == "Montaje general"
    assert leido.nodos[-1].hijos[0].descripcion == "MONTAJE DIETAS FICTICIAS"


def test_el_titulo_del_apartado_es_el_nombre_del_fichero():
    from app.services.presupuesto_excel import _preparar

    _, leido = _preparar(
        r"C:\obras\alimentacion-ficticia-v2.xlsm",
        _libro_formula(),
    )
    assert leido.nombre_apartado == "alimentacion-ficticia-v2"


def test_calcula_el_total_cuando_la_formula_no_trae_cache():
    leido = leer_excel_expertis(_libro_formula())

    assert leido.nombre_apartado == "Apartado"
    assert leido.partidas == 1
    assert leido.nodos[0].importe == Decimal("18.00")
    assert leido.nodos[0].descuento_pct == Decimal("10")
    assert leido.sumas_cuadran is True


def test_alta_desde_varios_excel_y_contraste_del_pdf(
    client: TestClient, db: Session, monkeypatch, tmp_path
):
    from app.config import get_settings

    monkeypatch.setenv("UPLOADS_DIR", str(tmp_path))
    get_settings.cache_clear()
    uno = _libro_descuadre()
    dos = _libro_formula()
    respuesta = client.post(
        "/proyectos/importar-presupuestos-excel",
        data={"nombre": "Obra ficticia norte"},
        files=[
            ("ficheros", ("ventilacion-v2.xlsx", uno, "application/vnd.ms-excel")),
            ("ficheros", ("paneles.xlsx", dos, "application/vnd.ms-excel")),
        ],
    )
    assert respuesta.status_code == 201, respuesta.text
    cuerpo = respuesta.json()
    assert cuerpo["requiere_revision"] is True
    assert cuerpo["sumas_cuadran"] is False
    assert cuerpo["partidas_detectadas"] == 4
    assert cuerpo["lineas_montaje"] == 1
    assert len(cuerpo["apartados"]) == 2

    proyecto_id = cuerpo["proyecto_id"]
    proyecto = db.get(Proyecto, proyecto_id)
    assert proyecto is not None
    assert proyecto.nombre == "Obra ficticia norte"
    assert proyecto.estado == "pendiente_revision"
    naves = db.query(Nave).filter(Nave.proyecto_id == proyecto_id).all()
    assert len(naves) == 1
    assert naves[0].importe_presupuestado == Decimal("73.00")

    tareas = (
        db.query(Tarea)
        .join(Nave, Tarea.nave_id == Nave.id)
        .filter(Nave.proyecto_id == proyecto_id, Tarea.es_montaje.is_(True))
        .all()
    )
    assert len(tareas) == 1
    assert {tarea.capitulo for tarea in tareas} == {"EXAFEED"}
    presupuestos = (
        db.query(Presupuesto).filter(Presupuesto.proyecto_id == proyecto_id).all()
    )
    assert {p.formato for p in presupuestos} == {"excel_expertis"}
    assert {p.version for p in presupuestos} == {2, 1}
    assert all(p.ruta_almacen for p in presupuestos)
    for presupuesto in presupuestos:
        descarga = client.get(
            f"/proyectos/{proyecto_id}/presupuestos/{presupuesto.id}/fichero"
        )
        assert descarga.status_code == 200, descarga.text
        assert descarga.content

    for tarea in tareas:
        tarea.capitulo = None
    db.commit()
    ventilacion = next(p for p in presupuestos if "ventilacion" in p.fichero_origen)
    from app.services.presupuesto_excel import rellenar_origen_montaje

    rellenar_origen_montaje(db, ventilacion)
    for tarea in tareas:
        db.refresh(tarea)
    assert {tarea.capitulo for tarea in tareas} == {"EXAFEED"}

    llamadas = {"n": 0}

    def extraer(_texto):
        llamadas["n"] += 1
        return PresupuestoExtraido(
            nombre_proyecto="Obra ficticia norte",
            naves=[
                NaveExtraida(
                    codigo="N1",
                    descripcion="Nave",
                    apartados=[
                        ApartadoExtraido(
                            codigo="S1",
                            descripcion="Salida de silos ficticia",
                            importe=Decimal("1000.00"),
                        ),
                        ApartadoExtraido(
                            codigo="M1",
                            descripcion="Cerramiento lateral ficticio",
                            importe=Decimal("35.00"),
                        ),
                    ],
                )
            ],
        )

    monkeypatch.setattr(
        "app.services.presupuesto_excel.extraccion_svc.extraer_presupuesto",
        extraer,
    )
    monkeypatch.setattr(
        "app.services.presupuesto_excel._relacionar_pdf_llm",
        lambda *_args, **_kwargs: None,
    )
    contraste = client.post(
        f"/proyectos/{proyecto_id}/contrastar-presupuesto-cliente",
        files={"fichero": ("contrato-cliente.md", b"Sistema de ventilacion ficticio", "text/markdown")},
    )
    assert contraste.status_code == 201, contraste.text
    sugerencia = contraste.json()["sugerencias"][0]
    silos = (
        db.query(Tarea)
        .join(Nave, Tarea.nave_id == Nave.id)
        .filter(
            Nave.proyecto_id == proyecto_id,
            Tarea.descripcion == "SALIDA DE SILOS",
            Tarea.nivel == "subapartado",
        )
        .one()
    )
    montaje = (
        db.query(Tarea)
        .join(Nave, Tarea.nave_id == Nave.id)
        .filter(
            Nave.proyecto_id == proyecto_id,
            Tarea.descripcion == "Montaje EXAFEED",
            Tarea.nivel == "subapartado",
        )
        .one()
    )
    assert sugerencia["tarea_id"] is None
    assert sugerencia["id"]
    assert sugerencia["importe_pdf"] == 1000.0
    repetido_pdf = client.post(
        f"/proyectos/{proyecto_id}/contrastar-presupuesto-cliente",
        files={"fichero": ("otra-copia.md", b"Sistema de ventilacion ficticio", "text/markdown")},
    )
    assert repetido_pdf.status_code == 422, repetido_pdf.text
    assert "ya está" in repetido_pdf.json()["detail"]
    assert llamadas["n"] == 1
    assert contraste.json()["sugerencias"][1]["tarea_id"] is None
    assert llamadas["n"] == 1

    refresco = client.get(f"/proyectos/{proyecto_id}/contraste-cliente")
    assert refresco.status_code == 200, refresco.text
    assert refresco.json()["presupuesto_id"] == contraste.json()["presupuesto_id"]
    assert llamadas["n"] == 1
    assert refresco.json()["sugerencias"][0]["id"] == sugerencia["id"]
    assert refresco.json()["sugerencias"][0]["confirmada"] is False

    detalle = client.get(f"/proyectos/{proyecto_id}")
    assert detalle.status_code == 200, detalle.text
    assert detalle.json()["contraste_cliente"]["sugerencias"][0]["descripcion_pdf"] == (
        "Salida de silos ficticia"
    )

    confirmado = client.post(
        f"/proyectos/{proyecto_id}/apartados/{silos.id}/importe-cliente",
        json={"importe": sugerencia["importe_pdf"]},
    )
    assert confirmado.status_code == 200, confirmado.text
    assert confirmado.json()["importe_contrato_cliente"] == 1000.0

    control = client.get(f"/proyectos/{proyecto_id}/control-economico")
    assert control.status_code == 200, control.text
    fila = next(
        item
        for item in control.json()["por_apartado"]
        if item["tarea_id"] == cuerpo["apartados"][0]["tarea_id"]
    )
    assert fila["presupuestado"] == 55.0
    assert fila["montaje"] == 30.0
    assert fila["contrato_cliente"] == 1000.0

    repartido = client.post(
        f"/proyectos/{proyecto_id}/importe-cliente",
        json={"tarea_ids": [silos.id, montaje.id], "importe": 100, "linea_id": sugerencia["id"]},
    )
    assert repartido.status_code == 200, repartido.text
    asignaciones = {
        item["tarea_id"]: item["importe_contrato_cliente"]
        for item in repartido.json()["asignaciones"]
    }
    assert abs(asignaciones[silos.id] - 40) < 0.001
    assert abs(asignaciones[montaje.id] - 60) < 0.001
    assert round(sum(asignaciones.values()), 2) == 100
    confirmadas = client.get(f"/proyectos/{proyecto_id}/contraste-cliente")
    assert confirmadas.status_code == 200, confirmadas.text
    assert llamadas["n"] == 1
    linea = next(
        fila for fila in confirmadas.json()["sugerencias"] if fila["id"] == sugerencia["id"]
    )
    assert linea["confirmada"] is True
    assert linea["sugerida"] is False
    assert set(linea["tarea_ids"]) == {silos.id, montaje.id}

    anadido = client.post(
        f"/proyectos/{proyecto_id}/presupuestos-excel",
        files={"fichero": ("otro.xlsx", _libro_subapartados(), "application/vnd.ms-excel")},
    )
    assert anadido.status_code == 201, anadido.text
    assert anadido.json()["apartados"][0]["codigo"] == "03"


def test_no_guarda_dos_veces_el_mismo_excel(client: TestClient, db: Session, monkeypatch, tmp_path):
    from app.config import get_settings
    from app.services.proyectos import borrar_proyecto

    monkeypatch.setenv("UPLOADS_DIR", str(tmp_path))
    get_settings.cache_clear()
    libro = _libro_formula()
    proyecto_id = None
    try:
        alta = client.post(
            "/proyectos/importar-presupuestos-excel",
            data={"nombre": "Obra ficticia sin repetir"},
            files=[("ficheros", ("alimentacion-ficticia.xlsx", libro, "application/vnd.ms-excel"))],
        )
        assert alta.status_code == 201, alta.text
        proyecto_id = alta.json()["proyecto_id"]
        otra_vez = client.post(
            f"/proyectos/{proyecto_id}/presupuestos-excel",
            files={"fichero": ("copia-ficticia.xlsx", libro, "application/vnd.ms-excel")},
        )
        assert otra_vez.status_code == 422, otra_vez.text
        assert "ya está" in otra_vez.json()["detail"]
        assert (
            db.query(Presupuesto)
            .filter(
                Presupuesto.proyecto_id == proyecto_id,
                Presupuesto.formato == "excel_expertis",
            )
            .count()
            == 1
        )
        lote = client.post(
            "/proyectos/importar-presupuestos-excel",
            data={"nombre": "Obra ficticia duplicada"},
            files=[
                ("ficheros", ("uno-ficticio.xlsx", libro, "application/vnd.ms-excel")),
                ("ficheros", ("dos-ficticio.xlsx", libro, "application/vnd.ms-excel")),
            ],
        )
        assert lote.status_code == 422, lote.text
        assert "mismo fichero" in lote.json()["detail"]
        assert db.query(Proyecto).filter(Proyecto.nombre == "Obra ficticia duplicada").count() == 0
    finally:
        if proyecto_id is not None:
            borrar_proyecto(db, proyecto_id)
        get_settings.cache_clear()


def test_el_descuento_de_cierre_se_reparte_y_el_manuscrito_manda():
    from app.schemas.extraccion import DescuentoExtraido, PresupuestoExtraido
    from app.schemas.presupuesto import SugerenciaContrasteOut
    from app.services.presupuesto_excel import _repartir_cierre

    lineas = [
        SugerenciaContrasteOut(
            descripcion_pdf="Silos ficticios",
            codigo_pdf="S",
            importe_pdf=Decimal("80.00"),
        ),
        SugerenciaContrasteOut(
            descripcion_pdf="Montaje ficticio",
            codigo_pdf="M",
            importe_pdf=Decimal("20.00"),
        ),
    ]
    extraido = PresupuestoExtraido(
        nombre_proyecto="Obra ficticia",
        total_impreso=Decimal("95.00"),
        total_manuscrito=Decimal("90.00"),
        descuentos=[
            DescuentoExtraido(descripcion="Descuento impreso", importe=Decimal("5.00")),
            DescuentoExtraido(
                descripcion="Descuento manuscrito",
                importe=Decimal("5.00"),
                es_manuscrito=True,
            ),
        ],
    )

    nuevas = _repartir_cierre(lineas, extraido)

    assert nuevas[0].importe_bruto == Decimal("80.00")
    assert nuevas[0].importe_pdf == Decimal("72.00")
    assert nuevas[1].importe_pdf == Decimal("18.00")
    assert sum(fila.importe_pdf for fila in nuevas) == Decimal("90.00")


def test_elementos_de_montaje_vuelve_a_su_subbloque(db: Session):
    from app.services.presupuesto_excel import devolver_elementos_de_montaje

    proyecto = Proyecto(nombre="Obra ficticia elementos", tipo="llave_en_mano", estado="borrador")
    db.add(proyecto)
    db.flush()
    nave = Nave(
        proyecto_id=proyecto.id,
        codigo="N1",
        descripcion="Nave ficticia",
        importe_presupuestado=Decimal("30.00"),
    )
    db.add(nave)
    db.flush()
    presupuesto = Presupuesto(
        proyecto_id=proyecto.id,
        version=1,
        fichero_origen="alimentacion-ficticia.xlsx",
        estado_extraccion="pendiente",
        es_escaneado=False,
        formato="excel_expertis",
    )
    db.add(presupuesto)
    db.flush()

    def tarea(**datos) -> Tarea:
        fila = Tarea(
            nave_id=nave.id,
            presupuesto_id=presupuesto.id,
            tiene_anotacion_manual=False,
            estado_revision="pendiente",
            estado="no_iniciada",
            avance_fisico_pct=Decimal("0"),
            **datos,
        )
        db.add(fila)
        db.flush()
        return fila

    apartado = tarea(
        codigo="01",
        nivel="apartado",
        descripcion="ALIMENTACION FICTICIA",
        importe_presupuestado=Decimal("30.00"),
    )
    montaje = tarea(
        tarea_padre_id=apartado.id,
        codigo="0",
        nivel="subapartado",
        descripcion="Montaje",
        importe_presupuestado=Decimal("15.00"),
    )
    elementos = tarea(
        tarea_padre_id=montaje.id,
        codigo="08",
        nivel="partida",
        descripcion="ELEMENTOS DE MONTAJE",
        capitulo="EXAFEED",
        importe_presupuestado=Decimal("5.00"),
        es_montaje=True,
        cantidad=Decimal("1"),
    )
    transportadores = tarea(
        tarea_padre_id=montaje.id,
        codigo="58",
        nivel="partida",
        descripcion="MONTAJE TRANSPORTADORES",
        capitulo="EXAFEED",
        importe_presupuestado=Decimal("10.00"),
        es_montaje=True,
        cantidad=Decimal("1"),
    )
    alimentacion = tarea(
        tarea_padre_id=apartado.id,
        codigo="08",
        nivel="subapartado",
        descripcion="ALIMENTACION",
        importe_presupuestado=Decimal("15.00"),
    )
    exafeed = tarea(
        tarea_padre_id=alimentacion.id,
        codigo="b5",
        nivel="subapartado",
        descripcion="EXAFEED",
        importe_presupuestado=Decimal("15.00"),
    )
    db.commit()

    try:
        devolver_elementos_de_montaje(db, presupuesto.id)
        db.refresh(elementos)
        db.refresh(transportadores)
        db.refresh(montaje)
        db.refresh(exafeed)
        db.refresh(alimentacion)
        db.refresh(apartado)

        assert elementos.tarea_padre_id == exafeed.id
        assert elementos.es_montaje is False
        assert elementos.capitulo is None
        assert transportadores.tarea_padre_id == montaje.id
        assert transportadores.es_montaje is True
        assert montaje.importe_presupuestado == Decimal("10.00")
        assert exafeed.importe_presupuestado == Decimal("20.00")
        assert alimentacion.importe_presupuestado == Decimal("20.00")
        assert apartado.importe_presupuestado == Decimal("30.00")

        from app.services.presupuesto_excel import repartir_montaje_por_subapartado

        repartir_montaje_por_subapartado(db, presupuesto.id)
        db.refresh(transportadores)
        db.refresh(exafeed)
        db.refresh(alimentacion)
        db.refresh(elementos)
        bloque = db.get(Tarea, transportadores.tarea_padre_id)
        assert bloque is not None
        assert bloque.descripcion == "Montaje EXAFEED"
        assert bloque.tarea_padre_id == exafeed.id
        assert transportadores.es_montaje is True
        assert bloque.importe_presupuestado == Decimal("10.00")
        assert exafeed.importe_presupuestado == Decimal("30.00")
        assert alimentacion.importe_presupuestado == Decimal("30.00")
        assert apartado.importe_presupuestado == Decimal("30.00")
        assert db.get(Tarea, montaje.id) is None
        assert elementos.tarea_padre_id == exafeed.id

        devolver_elementos_de_montaje(db, presupuesto.id)
        db.refresh(elementos)
        assert elementos.tarea_padre_id == exafeed.id
    finally:
        db.query(Tarea).filter(Tarea.presupuesto_id == presupuesto.id).delete()
        db.delete(presupuesto)
        db.delete(nave)
        db.delete(proyecto)
        db.commit()


def test_no_contrasta_un_bloque_y_sus_subapartados(db: Session):
    from app.services.presupuesto_excel import confirmar_importes_cliente

    proyecto = Proyecto(nombre="Obra ficticia contraste", tipo="llave_en_mano", estado="borrador")
    db.add(proyecto)
    db.flush()
    nave = Nave(
        proyecto_id=proyecto.id,
        codigo="N1",
        descripcion="Nave ficticia",
        importe_presupuestado=Decimal("30.00"),
    )
    db.add(nave)
    db.flush()
    presupuesto = Presupuesto(
        proyecto_id=proyecto.id,
        version=1,
        fichero_origen="ventilacion-ficticia.xlsx",
        estado_extraccion="pendiente",
        es_escaneado=False,
        formato="excel_expertis",
    )
    db.add(presupuesto)
    db.flush()

    def tarea(**datos) -> Tarea:
        fila = Tarea(
            nave_id=nave.id,
            presupuesto_id=presupuesto.id,
            tiene_anotacion_manual=False,
            estado_revision="pendiente",
            estado="no_iniciada",
            avance_fisico_pct=Decimal("0"),
            **datos,
        )
        db.add(fila)
        db.flush()
        return fila

    apartado = tarea(
        codigo="01",
        nivel="apartado",
        descripcion="VENTILACION FICTICIA",
        importe_presupuestado=Decimal("30.00"),
    )
    bloque = tarea(
        tarea_padre_id=apartado.id,
        codigo="07",
        nivel="subapartado",
        descripcion="VENTANAS FICTICIAS",
        importe_presupuestado=Decimal("30.00"),
    )
    hojas = tarea(
        tarea_padre_id=bloque.id,
        codigo="b1",
        nivel="subapartado",
        descripcion="HOJAS FICTICIAS",
        importe_presupuestado=Decimal("20.00"),
    )
    marco = tarea(
        tarea_padre_id=bloque.id,
        codigo="b2",
        nivel="subapartado",
        descripcion="MARCO FICTICIO",
        importe_presupuestado=Decimal("10.00"),
    )
    db.commit()

    try:
        with pytest.raises(ValueError, match="subapartados"):
            confirmar_importes_cliente(db, proyecto.id, [bloque.id, hojas.id], Decimal("30"))
        with pytest.raises(ValueError, match="subapartados"):
            confirmar_importes_cliente(
                db, proyecto.id, [bloque.id, marco.id], Decimal("30")
            )

        guardado = confirmar_importes_cliente(
            db, proyecto.id, [hojas.id, marco.id], Decimal("30")
        )
        por_id = {
            fila.tarea_id: fila.importe_contrato_cliente for fila in guardado.asignaciones
        }
        assert por_id[hojas.id] == Decimal("20.00")
        assert por_id[marco.id] == Decimal("10.00")
        db.refresh(bloque)
        assert bloque.importe_contrato_cliente is None
    finally:
        db.query(Tarea).filter(Tarea.presupuesto_id == presupuesto.id).delete()
        db.delete(presupuesto)
        db.delete(nave)
        db.delete(proyecto)
        db.commit()


def test_la_propuesta_marca_el_bloque_sin_confirmar(db: Session, monkeypatch):
    from app.models import LineaPdfCliente
    from app.services.presupuesto_excel import (
        _salida_contraste,
        _proponer_relaciones,
        confirmar_importes_cliente,
    )
    from app.services.proyectos import borrar_proyecto

    proyecto = Proyecto(nombre="Obra ficticia propuesta", tipo="llave_en_mano", estado="borrador")
    db.add(proyecto)
    db.flush()
    nave = Nave(
        proyecto_id=proyecto.id,
        codigo="N1",
        descripcion="Nave ficticia",
        importe_presupuestado=Decimal("30.00"),
    )
    db.add(nave)
    db.flush()
    excel = Presupuesto(
        proyecto_id=proyecto.id,
        version=1,
        fichero_origen="ventilacion-ficticia.xlsx",
        estado_extraccion="pendiente",
        es_escaneado=False,
        formato="excel_expertis",
    )
    pdf = Presupuesto(
        proyecto_id=proyecto.id,
        version=1,
        fichero_origen="contrato-ficticio.md",
        estado_extraccion="extraido",
        es_escaneado=False,
        formato="pdf_cliente",
    )
    db.add_all([excel, pdf])
    db.flush()

    def tarea(presupuesto_id: int, **datos) -> Tarea:
        fila = Tarea(
            nave_id=nave.id,
            presupuesto_id=presupuesto_id,
            tiene_anotacion_manual=False,
            estado_revision="pendiente",
            estado="no_iniciada",
            avance_fisico_pct=Decimal("0"),
            **datos,
        )
        db.add(fila)
        db.flush()
        return fila

    apartado = tarea(
        excel.id,
        codigo="01",
        nivel="apartado",
        descripcion="VENTILACION FICTICIA",
        importe_presupuestado=Decimal("30.00"),
    )
    bloque = tarea(
        excel.id,
        tarea_padre_id=apartado.id,
        codigo="07",
        nivel="subapartado",
        descripcion="VENTANAS FICTICIAS",
        importe_presupuestado=Decimal("30.00"),
    )
    hojas = tarea(
        excel.id,
        tarea_padre_id=bloque.id,
        codigo="b1",
        nivel="subapartado",
        descripcion="HOJAS FICTICIAS",
        importe_presupuestado=Decimal("20.00"),
    )
    tarea(
        excel.id,
        tarea_padre_id=bloque.id,
        codigo="b2",
        nivel="subapartado",
        descripcion="MARCO FICTICIO",
        importe_presupuestado=Decimal("10.00"),
    )
    linea = LineaPdfCliente(
        presupuesto_id=pdf.id,
        orden=0,
        tipo="importe",
        codigo="V1",
        descripcion="Ventanas ficticias",
        importe=Decimal("30.00"),
    )
    db.add(linea)
    db.commit()

    llamadas = {"n": 0}

    def relacion(_lineas, _bloques):
        llamadas["n"] += 1
        return {0: [0, 1]}

    monkeypatch.setattr("app.services.presupuesto_excel._relacionar_pdf_llm", relacion)
    try:
        assert _proponer_relaciones(db, proyecto.id, pdf.id) is True
        db.commit()
        assert llamadas["n"] == 1
        salida = _salida_contraste(db, pdf)
        fila = salida.sugerencias[0]
        assert fila.confirmada is False
        assert fila.sugerida is True
        assert fila.tarea_ids == [bloque.id]
        assert _proponer_relaciones(db, proyecto.id, pdf.id) is False
        assert llamadas["n"] == 1

        confirmar_importes_cliente(db, proyecto.id, [hojas.id], Decimal("30.00"), linea.id)
        cerrada = _salida_contraste(db, pdf).sugerencias[0]
        assert cerrada.confirmada is True
        assert cerrada.sugerida is False
        assert cerrada.tarea_ids == [hojas.id]
    finally:
        borrar_proyecto(db, proyecto.id)


def test_el_nombre_elige_la_parte_y_no_repite_el_bloque():
    from app.services.presupuesto_excel import (
        _BloqueContraste,
        _depurar_asignaciones,
        _interpretar_relaciones,
        _relacionar_por_nombre,
    )

    bloques = [
        _BloqueContraste(1, None, "Ventilacion ficticia / Ventanas ficticias"),
        _BloqueContraste(2, 1, "Ventilacion ficticia / Ventanas ficticias / Hojas ficticias"),
        _BloqueContraste(3, 1, "Ventilacion ficticia / Ventanas ficticias / Marco ficticio"),
    ]
    assert _relacionar_por_nombre(["Ventanas ficticias"], bloques) == {0: [0]}
    assert _relacionar_por_nombre(["Hojas ficticias", "Marco ficticio"], bloques) == {0: [1], 1: [2]}
    assert _depurar_asignaciones({0: [0, 1]}, ["Ventanas ficticias"], bloques) == {0: [1]}
    assert _depurar_asignaciones(
        {0: [1], 1: [1]},
        ["Hojas ficticias", "Marco ficticio"],
        bloques,
    ) == {0: [2]}
    assert _interpretar_relaciones(
        [{"linea": 0, "bloques": [0, 5, "1"]}, {"linea": 3, "bloques": [0]}],
        1,
        2,
    ) == {0: [0, 1]}
